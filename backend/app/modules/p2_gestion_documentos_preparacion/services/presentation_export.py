"""Editable, bounded PPTX export. No inference or external network requests."""
import logging
import warnings
from io import BytesIO

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from pypdf import PdfReader

logger = logging.getLogger(__name__)


def extract_pdf_figures(content: bytes) -> list[tuple[int, bytes]]:
    """Up to three raster figures from the first ten pages, each <=8M pixels."""
    figures: list[tuple[int, bytes]] = []
    if not content.startswith(b"%PDF-"):
        return figures
    try:
        reader = PdfReader(BytesIO(content))
        for page_number, page in enumerate(reader.pages[:10], 1):
            resources = page.get("/Resources", {})
            if hasattr(resources, "get_object"):
                resources = resources.get_object()
            objects = resources.get("/XObject", {})
            if hasattr(objects, "get_object"):
                objects = objects.get_object()
            for name in list(page.images.keys())[:8]:
                # Skip nested/vector graphics rather than guess their representation.
                if not isinstance(name, str) or name not in objects:
                    continue
                raw = objects[name].get_object()
                width, height = int(raw.get("/Width", 0)), int(raw.get("/Height", 0))
                if min(width, height) < 120 or width * height > 8_000_000:
                    continue
                try:
                    with warnings.catch_warnings():
                        warnings.simplefilter("error", Image.DecompressionBombWarning)
                        image = page.images[name].image
                        image.thumbnail((1280, 720))
                        result = BytesIO()
                        image.convert("RGB").save(result, format="JPEG", quality=85)
                        figures.append((page_number, result.getvalue()))
                    if len(figures) == 3:
                        return figures
                except Exception:
                    logger.info("Skipping unsupported PDF image on page %s", page_number)
    except Exception:
        logger.info("PDF figures unavailable; exporting the presentation text")
    return figures


def export_presentation(material, *, source_name: str, pdf_content: bytes | None = None) -> bytes:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.title = material.title or "Defensa académica"
    deck.core_properties.author = "Socratia"

    def text(slide, value, x, y, width, height, size=24, color="203A43", bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        box.text_frame.word_wrap = True
        box.text_frame.margin_left = box.text_frame.margin_right = Inches(.03)
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = value
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        return box

    def new_slide(title, position):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F5F8F7")
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, deck.slide_width, Inches(.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string("278B78")
        bar.line.fill.background()
        text(slide, title, .65, .45, 12, 1.15, 30, bold=True)
        text(slide, f"Socratia · {source_name[:95]} · {position:02}", .65, 7.02, 12, .3, 10)
        return slide

    for item in material.slides:
        slide = new_slide(item.title, item.position)
        objective = item.objective if len(item.objective) <= 220 else item.objective[:217] + "…"
        text(slide, objective, .7, 1.65, 11.9, .85, 17, color="526E74")
        bullets = item.bullet_points[:5]
        for index, bullet in enumerate(bullets):
            # A long original bullet stays complete in the notes, while the slide stays readable.
            visible = bullet if len(bullet) <= 190 else bullet[:187] + "…"
            text(slide, f"•  {visible}", .9, 2.65 + index * .78, 11.5, .75, 20 if len(visible) > 145 else 24)
        notes = slide.notes_slide.notes_text_frame
        if notes is not None:
            notes.text = (f"Objetivo: {item.objective}\n\n{item.speaker_notes}\n\n"
                          f"Tiempo estimado: {item.estimated_seconds} segundos\n\n"
                          + "\n".join(item.bullet_points))

    figures = extract_pdf_figures(pdf_content) if pdf_content else []
    for page_number, image_bytes in figures:
        slide = new_slide(f"Figura del documento · página {page_number}", len(deck.slides) + 1)
        with Image.open(BytesIO(image_bytes)) as image:
            width, height = image.size
        scale = min(11.5 / width, 4.6 / height)
        slide.shapes.add_picture(BytesIO(image_bytes), Inches((13.333 - width * scale) / 2),
                                 Inches(1.8), width=Inches(width * scale), height=Inches(height * scale))
        text(slide, "Anexo visual del PDF. Revisa su relación con tu exposición antes de presentarlo.",
             .7, 6.5, 11.9, .45, 14)

    output = BytesIO()
    deck.save(output)
    return output.getvalue()
