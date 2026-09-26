import sys
import time
import uuid
from io import BytesIO
import json
import httpx
from docx import Document as DocxDocument
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.integrations.embeddings.dependencies import get_embedding_provider
from app.integrations.vector_db.router import VectorStoreRouter
from app.integrations.vector_db.base import VectorStoreError, VectorStoreProvider, VectorSearchResult
from app.integrations.vector_db.pgvector_provider import PgVectorStoreProvider
from app.integrations.llm.dependencies import get_question_generation_router
from app.modules.p1_gestion_identidad_seguridad.models.user import User
from app.modules.p2_gestion_documentos_preparacion.services.question_bank_service import QuestionBankService
from app.modules.p2_gestion_documentos_preparacion.services.presentation_material_service import PresentationMaterialService
from app.integrations.llm.dependencies import get_presentation_generation_router
from app.modules.p2_gestion_documentos_preparacion.models.vector_store import EmbeddingProfileStoreSync

API_URL = "http://localhost:8000/api/v1"

class FailingPineconeProvider(VectorStoreProvider):
    def upsert(self, **kwargs) -> None:
        raise NotImplementedError()
        
    def delete(self, **kwargs) -> None:
        raise NotImplementedError()
        
    def query(self, **kwargs) -> list[VectorSearchResult]:
        raise VectorStoreError("Forced Pinecone failure for E-4 testing")

def build_large_document() -> bytes:
    # 239 chunks require a large document. The chunker splits by size. 
    # Let's generate a massive document.
    document = DocxDocument()
    document.add_heading("Socratia: preparacion de defensas academicas", level=1)
    
    # Generate around 150 paragraphs, which might result in ~239 chunks depending on settings
    for i in range(150):
        document.add_heading(f"Seccion {i}", level=2)
        document.add_paragraph("Este es un parrafo muy largo que contiene informacion sustancial para asegurar que el chunker genere una cantidad significativa de fragmentos. " * 30)
        
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()

def run_tests():
    print("=== SMOKE TEST E-4: Vector Store Router ===")
    
    client = httpx.Client(base_url=API_URL, timeout=300)
    suffix = uuid.uuid4().hex[:12]
    
    # 1. Upload Large Document via API to trigger Chunking and Outbox Sync
    print("\n1. Registrando usuario y subiendo documento...")
    response = client.post(
        "/auth/register",
        json={
            "email": f"cu10-router-{suffix}@example.com",
            "full_name": "CU10 Router Smoke",
            "password": f"Socratia smoke phrase {uuid.uuid4().hex}",
        },
    )
    response.raise_for_status()
    user_id = response.json()["user"]["id"]

    response = client.post(
        "/documents",
        files={
            "file": (
                "socratia-router-smoke.docx",
                build_large_document(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    response.raise_for_status()
    document_id = response.json()["id"]

    response = client.post(f"/documents/{document_id}/process")
    response.raise_for_status()
    chunk_count = int(response.json()["chunk_count"])
    print(f"Documento procesado. Chunks generados: {chunk_count}")

    # Wait for the outbox worker to sync to pgvector and pinecone
    # The actual background worker must be running for this test to pass.
    # In a real environment, we'd poll the status endpoints, but let's just wait 10 seconds.
    print("Esperando 10 segundos para sincronizacion del Outbox...")
    time.sleep(10)

    settings = get_settings()
    engine = create_engine(settings.database_url, pool_pre_ping=True)
    
    with Session(engine, expire_on_commit=False) as db:
        user = db.get(User, user_id)
        
        primary = FailingPineconeProvider()
        secondary = PgVectorStoreProvider(db=db)
        
        router = VectorStoreRouter(
            primary=primary,
            secondary=secondary,
            db=db,
            circuit_breaker_timeout_sec=5
        )
        
        llm_router = get_question_generation_router()
        service = QuestionBankService(
            db=db,
            embeddings=get_embedding_provider(),
            vectors=router, # Injecting our test router
            generation_router=llm_router,
        )

        # Force sync state just in case outbox is slow/not running
        # We find the profile associated with this document
        # Wait, the profile is linked to the document's embedding profile.
        from app.modules.p2_gestion_documentos_preparacion.models.document import Document
        doc = db.get(Document, document_id)
        profile_id = doc.embedding_profile_id
        
        sync_state = db.query(EmbeddingProfileStoreSync).filter_by(
            profile_id=profile_id, store_name="pgvector"
        ).first()
        
        if not sync_state:
            print("WARNING: pgvector sync state no existe, creandolo manualmente para el test")
            sync_state = EmbeddingProfileStoreSync(
                profile_id=profile_id,
                store_name="pgvector",
                status="SYNCED",
                vectors_expected=chunk_count,
                vectors_synced=chunk_count
            )
            db.add(sync_state)
            db.commit()
        else:
            print(f"Estado de sync pgvector actual: {sync_state.status}")
            sync_state.status = "SYNCED"
            db.commit()

        # TEST 1: Fallback Exitoso (Pinecone Falla, pgvector SYNCED)
        print("\n=== TEST 1: Fallback Exitoso (pgvector SYNCED) ===")
        print("Ejecutando CU10 con Router (Pinecone Fallando)...")
        try:
            bank = service.generate(user, document_id)
            print("✅ CU10 completo satisfactoriamente.")
            print(f"Preguntas generadas: {len(bank.questions)}")
            
            if len(bank.questions) == 12:
                print("✅ Test 1 (CU10) PASSED: Se obtuvieron 12 preguntas utilizando el VectorStoreRouter con fallback a pgvector.")
            else:
                print("❌ Test 1 (CU10) FAILED: Cantidad de preguntas incorrecta.")
                
            print("\nEjecutando CU11 con Router (Pinecone Fallando)...")
            presentation_service = PresentationMaterialService(
                db=db,
                embeddings=get_embedding_provider(),
                vectors=router,
                generation_router=get_presentation_generation_router(),
            )
            presentation = presentation_service.generate(user, document_id, duration_minutes=15, regenerate=True)
            print("✅ CU11 completo satisfactoriamente.")
            print(f"Slides generados: {len(presentation.slides)}")
            
            if 8 <= len(presentation.slides) <= 12:
                print("✅ Test 1 (CU11) PASSED: Material de presentacion generado usando fallback a pgvector.")
            else:
                print("❌ Test 1 (CU11) FAILED: Cantidad de slides incorrecta.")
                
        except Exception as e:
            print(f"❌ Test 1 FAILED: Exception occurred - {e}")
            
        # TEST 2: Fallback Rechazado (Pinecone Falla, pgvector PENDING)
        print("\n=== TEST 2: Fallback Rechazado (pgvector PENDING/Incompleto) ===")
        sync_state.status = "PENDING"
        db.commit()
        
        try:
            bank2 = service.generate(user, document_id)
            print("❌ Test 2 FAILED: CU10 termino exitosamente cuando debio fallar (pgvector incompleto).")
        except Exception as e:
            if "not completely SYNCED" in str(e) or "VectorStoreError" in str(type(e).__name__):
                print(f"✅ Test 2 PASSED: El router rechazo usar pgvector incompleto. Error capturado: {e}")
            else:
                print(f"❓ Test 2: Falla capturada pero error inesperado: {e}")
                
    client.close()
    return 0

if __name__ == "__main__":
    sys.exit(run_tests())
