from io import BytesIO
from types import SimpleNamespace

from PIL import Image
from pptx import Presentation

from app.modules.p2_gestion_documentos_preparacion.services.presentation_export import (
    export_presentation,
    extract_pdf_figures,
)


def sample_material():
    return SimpleNamespace(title="Defensa de prueba", slides=[SimpleNamespace(
        title="Resultados", position=1, objective="Explicar la evidencia obtenida",
        bullet_points=["Resultado documentado", "Limitación conocida"],
        speaker_notes="Explicar el método y sus limitaciones sin inventar resultados.",
        estimated_seconds=90,
    )])


def test_powerpoint_has_editable_text_and_speaker_notes():
    deck = Presentation(BytesIO(export_presentation(sample_material(), source_name="tesis.docx")))
    assert len(deck.slides) == 1
    contents = "\n".join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)
    assert "Resultados" in contents and "Resultado documentado" in contents
    assert "90 segundos" in deck.slides[0].notes_slide.notes_text_frame.text


def test_pdf_image_is_an_attributed_appendix_not_fabricated_slide_evidence():
    pdf = BytesIO()
    Image.new("RGB", (320, 200), "navy").save(pdf, format="PDF")
    figures = extract_pdf_figures(pdf.getvalue())
    assert len(figures) == 1 and figures[0][0] == 1
    deck = Presentation(BytesIO(export_presentation(sample_material(), source_name="tesis.pdf", pdf_content=pdf.getvalue())))
    assert len(deck.slides) == 2
    assert any(shape.shape_type == 13 for shape in deck.slides[1].shapes)
    assert any("página 1" in shape.text for shape in deck.slides[1].shapes if shape.has_text_frame)


def test_unsupported_pdf_does_not_break_text_export():
    assert extract_pdf_figures(b"not a PDF") == []
    deck = Presentation(BytesIO(export_presentation(sample_material(), source_name="tesis.pdf", pdf_content=b"%PDF-broken")))
    assert len(deck.slides) == 1


def test_export_requires_authentication(client):
    assert client.get("/api/v1/documents/missing/presentation/export").status_code == 401
