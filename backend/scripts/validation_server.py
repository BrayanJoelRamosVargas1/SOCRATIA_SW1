"""Isolated UI validation backend. External providers use deterministic test doubles.
Never point this runner at a development/production database.
"""

import argparse
import json
import os
import sys
from pathlib import Path
from types import SimpleNamespace

backend = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(backend))
sys.path.insert(0, str(backend / "tests"))
parser = argparse.ArgumentParser()
parser.add_argument("--runtime", type=Path, required=True)
args = parser.parse_args()
config = json.loads(args.runtime.read_text(encoding="utf-8-sig"))
os.environ["DATABASE_URL"] = (
    f"postgresql+psycopg://socratia:{config['password']}@127.0.0.1:{config['port']}/socratia_ui_test"
)
os.environ["CORS_ORIGINS"] = "http://localhost:3000,http://127.0.0.1:3000"
os.environ["COOKIE_SECURE"] = "false"
os.environ["LOCAL_STORAGE_PATH"] = str(Path(config["root"]) / "ui-uploads")

import conftest as fixtures
from fastapi import Depends
from sqlalchemy.orm import Session

from app.core.database import Base, SessionLocal, engine, get_db
from app.core.security import hash_password
from app.modules.p1_gestion_identidad_seguridad.repositories.user_repository import UserRepository
from app.modules.p4_gestion_evaluacion_resultados.schemas.evaluation import TurnEvaluationResult
from app.modules.p4_gestion_evaluacion_resultados.services import evaluation_service

if engine.url.database != "socratia_ui_test" or engine.url.port != config["port"]:
    raise RuntimeError("Validation requires its isolated UI test database")
# The native Windows test cluster does not ship pgvector. This domain only
# exercises serialization; FakeVectorStoreProvider replaces similarity search.
from sqlalchemy import text

with engine.begin() as connection:
    if not connection.scalar(text("SELECT 1 FROM pg_type WHERE typname='vector'")):
        connection.execute(text("CREATE DOMAIN vector AS text"))
Base.metadata.create_all(engine)
app = fixtures.app
embed = fixtures.FakeEmbeddingProvider()
namespace_records = {}


def vectors(db: Session = Depends(get_db)):
    result = fixtures.FakeVectorStoreProvider(db)
    result._namespaces = namespace_records
    return result


app.dependency_overrides[fixtures.get_embedding_provider] = lambda: embed
app.dependency_overrides[fixtures.get_vector_store_provider] = vectors
app.dependency_overrides[fixtures.get_email_provider] = lambda: fixtures.FakeEmailProvider()
app.dependency_overrides[fixtures.get_question_generation_router] = lambda: (
    fixtures.QuestionGenerationRouter(
        primary=fixtures.FakeQuestionGenerationProvider(name="validation", model="fixture"),
        fallback=fixtures.FakeQuestionGenerationProvider(
            name="validation-fallback", model="fixture"
        ),
        failure_threshold=3,
        recovery_seconds=60,
    )
)
app.dependency_overrides[fixtures.get_presentation_generation_router] = lambda: (
    fixtures.PresentationGenerationRouter(
        primary=fixtures.FakePresentationGenerationProvider(name="validation", model="fixture"),
        fallback=fixtures.FakePresentationGenerationProvider(
            name="validation-fallback", model="fixture"
        ),
        failure_threshold=3,
        recovery_seconds=60,
    )
)

# Keep accidental microphone input inside this validation process. Never call
# real STT, routing or paid speech providers from the disposable UI pilot.
from app.integrations.llm.jury_base import JuryAction
from app.integrations.tts.base import TTSProviderError
from app.modules.p3_gestion_simulacion.services import voice_service


class ValidationSTT:
    primary = SimpleNamespace(name="validation")

    async def transcribe(self, request):
        return "Respuesta sintética del entorno de validación; no es una transcripción real.", False, 0


class ValidationJury:
    async def route(self, request):
        return SimpleNamespace(decision=SimpleNamespace(action=JuryAction.NEXT_BANK, follow_up_question=None))


class ValidationTTS:
    def synthesize(self, request):
        raise TTSProviderError("UI validation uses browser speech", provider="validation")


voice_service.get_stt_router = lambda: ValidationSTT()
voice_service.get_jury_routing_router = lambda **kwargs: ValidationJury()
voice_service.get_tts_provider = lambda: ValidationTTS()


class ValidationEvaluator:
    provider = "validation"
    model = "fixture"

    def __init__(self, **kwargs):
        pass

    async def evaluate_turn(self, question, student_response, context_text, expected_points=None):
        return TurnEvaluationResult(
            classification="PRECISE",
            score=82,
            strengths=["Respuesta registrada y vinculada al documento de prueba."],
            weaknesses=["Amplía la evidencia en tu próximo intento."],
            missing_points=[],
            contradictions=[],
            evidence=[student_response[:160]],
        )


evaluation_service.Evaluator = ValidationEvaluator
with SessionLocal() as db:
    repo = UserRepository(db)
    for email, name, role in [
        ("validation.admin@example.com", "Administración de prueba", "Administrator"),
        ("validation.student@example.com", "Estudiante de prueba", "student"),
    ]:
        if repo.get_by_email(email) is None:
            user = repo.create(
                email=email,
                full_name=name,
                password_hash=hash_password("Socratia-validacion-2026!"),
            )
            user.roles.append(repo.get_or_create_role(role))
    db.commit()
from docx import Document

sample = Document()
sample.add_heading("Defensa de validación Socratia", 0)
sample.add_paragraph(
    (
        "La investigación evalúa un sistema de preparación de defensas. El objetivo es mejorar la argumentación mediante preguntas contextualizadas. La metodología combina revisión documental y ensayos controlados. Los resultados se contrastan con la evidencia y se reconocen las limitaciones. "
    )
    * 80
)
sample.save(Path(config["root"]) / "tesis-validacion.docx")
print(
    "VALIDATION_BACKEND_READY: isolated PostgreSQL and deterministic external providers", flush=True
)
import uvicorn

uvicorn.run(app, host="127.0.0.1", port=8000, log_level="warning")
