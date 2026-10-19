"""One bounded cloud document/presentation journey using a fictional PDF."""
from io import BytesIO
import json
import os
from pathlib import Path
import textwrap
import time

import httpx
from PIL import Image, ImageDraw
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

BASE = 'https://socratia.34.204.9.87.sslip.io'

def synthetic_pdf():
    writer=PdfWriter()
    sections=[
        ('Problema y objetivo','Proyecto universitario ficticio de agenda de tutorias. El problema es la coordinacion manual y la perdida de solicitudes. El objetivo es reducir duplicados y registrar la asistencia. No se usaron datos reales de alumnos.'),
        ('Metodologia','Se desarrollo un prototipo con arquitectura cliente servidor, API y PostgreSQL. Se trabajaron dos iteraciones, pruebas funcionales y entrevistas simuladas. La muestra ficticia incluye veinte participantes; no representa a toda la universidad.'),
        ('Implementacion','El prototipo permite solicitar una tutoria, verificar disponibilidad y confirmar la reserva. Un identificador unico evita reservas duplicadas. La autenticacion separa alumno y administrador. Las claves se almacenan fuera del codigo.'),
        ('Evaluacion','Se ejecutaron treinta casos funcionales del prototipo ficticio. Veintiocho pasaron y dos detectaron problemas de calendario. Se corrigieron los defectos y se repitieron los casos. El tiempo mediano simulado bajo de doce a cuatro minutos; son valores de prueba, no resultados reales.'),
        ('Limitaciones y conclusiones','La evaluacion es pequena, ficticia y no permite afirmar causalidad ni capacidad masiva. Falta medir accesibilidad, carga concurrente y usuarios reales. El trabajo futuro incluye una prueba controlada y una politica de respaldos. El prototipo muestra viabilidad tecnica limitada al alcance ensayado.'),
    ]
    font=writer._add_object(DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')}))
    page=writer.add_blank_page(width=595,height=842)
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
    lines=['SOCRATIA - DOCUMENTO FICTICIO DE VALIDACION','Agenda de tutorias universitarias','']
    for title,body in sections:
        lines += [title]+textwrap.wrap(body,width=82)+['']
    escape=lambda value:value.replace('\\','\\\\').replace('(','\\(').replace(')','\\)')
    stream=DecodedStreamObject()
    stream.set_data(('BT /F1 10 Tf 45 790 Td 13 TL '+' '.join('('+escape(line)+') Tj T*' for line in lines)+' ET').encode('ascii'))
    page[NameObject('/Contents')]=writer._add_object(stream)
    image=Image.new('RGB',(1000,480),'#f4f7f6');draw=ImageDraw.Draw(image)
    for x,label in [(40,'SOLICITUD'),(365,'DISPONIBILIDAD'),(690,'CONFIRMACION')]:
        draw.rounded_rectangle((x,180,x+270,300),radius=20,fill='#185d64')
        draw.text((x+30,230),label,fill='white')
    draw.line((310,240,365,240),fill='#278b78',width=8)
    draw.line((635,240,690,240),fill='#278b78',width=8)
    figure=BytesIO();image.save(figure,format='PDF');figure.seek(0);writer.append(figure)
    output=BytesIO();writer.write(output);return output.getvalue()


def main():
    email=os.environ['SOCRATIA_TRIAL_EMAIL'];password=os.environ['SOCRATIA_TRIAL_PASSWORD']
    root=Path(__file__).resolve().parents[2]/'artifacts/defense-validation';root.mkdir(exist_ok=True,parents=True)
    report={'origin':BASE,'synthetic_document':True,'stages':[]}
    with httpx.Client(base_url=BASE,timeout=90) as client:
        def call(stage,method,path,**kwargs):
            started=time.perf_counter();response=client.request(method,path,**kwargs)
            item={'stage':stage,'status':response.status_code,'elapsed_ms':round((time.perf_counter()-started)*1000)}
            report['stages'].append(item)
            print(json.dumps(item),flush=True)
            response.raise_for_status();return response
        try:
            call('login','POST','/api/v1/auth/login',json={'email':email,'password':password})
            pdf=synthetic_pdf();(root/'tesis-ficticia-rendimiento.pdf').write_bytes(pdf)
            doc=call('upload','POST','/api/v1/documents',files={'file':('tesis-ficticia-rendimiento.pdf',pdf,'application/pdf')}).json()
            report['document_id']=doc['id'];prefix=f'/api/v1/documents/{doc["id"]}'
            call('process','POST',prefix+'/process')
            material=call('presentation','POST',prefix+'/presentation/generate',json={'duration_minutes':5}).json()
            report['slides']=len(material['slides'])
            deck=call('export','GET',prefix+'/presentation/export').content
            (root/'defensa-ficticia-nube.pptx').write_bytes(deck)
            from pptx import Presentation
            parsed=Presentation(BytesIO(deck))
            report['exported_slides']=len(parsed.slides)
            report['embedded_figures']=sum(shape.shape_type==13 for slide in parsed.slides for shape in slide.shapes)
            report['ok']=report['exported_slides']>=report['slides'] and report['embedded_figures']>0
        except Exception as error:
            report['ok']=False;report['error_type']=type(error).__name__
        finally:
            (root/'cloud-document-flow.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps(report),flush=True)
    return 0 if report['ok'] else 1

if __name__=='__main__':raise SystemExit(main())
